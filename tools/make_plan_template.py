#!/usr/bin/env python3
"""
Generate the AutoDock planning workbook.

This is a dev-time tool, not part of the runtime (the server itself still has
zero dependencies). It reads the live database read-only so the dropdowns carry
the real workstream codes and the real team roster, then writes a workbook the
manager fills in by hand.

    python3 tools/make_plan_template.py

Re-run it whenever workstreams or team members change.
"""
import os
import re
import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "templates" / "AutoDock-Plan-Template.xlsx"

# --- look at what actually exists ------------------------------------------


def read_env(name, default):
    env = ROOT / ".env"
    if env.exists():
        for line in env.read_text().splitlines():
            m = re.match(r"\s*%s\s*=\s*(.*)" % re.escape(name), line)
            if m:
                return m.group(1).strip()
    return default


def load_context():
    db_path = read_env("DATABASE_PATH", "./data/team-ops.db")
    path = (ROOT / db_path.lstrip("./")).resolve()
    if not path.exists():
        print(f"warning: no database at {path}; using empty reference lists", file=sys.stderr)
        return {"project": None, "workstreams": [], "team": []}
    con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        project = con.execute("SELECT code, name FROM projects ORDER BY code LIMIT 1").fetchone()
        workstreams = con.execute(
            "SELECT code, name FROM workstreams ORDER BY code"
        ).fetchall()
        team = con.execute(
            "SELECT name, email, role, COALESCE(title, '') FROM users WHERE active = 1 ORDER BY name"
        ).fetchall()
    finally:
        con.close()
    return {"project": project, "workstreams": workstreams, "team": team}


MAX_SPAN = int(read_env("MAX_TASK_SPAN_DAYS", "3") or 3)
CTX = load_context()
WS_CODES = [c for c, _ in CTX["workstreams"]]

# --- styling ---------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
REQ_FILL = PatternFill("solid", fgColor="C00000")
EXAMPLE_FONT = Font(color="808080", italic=True)
EXAMPLE_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LAST_ROW = 400  # how far down validation and formatting reach


def write_sheet(wb, title, columns, example, note):
    """columns: list of (header, width, required, help_text, number_format)."""
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = "1F3864"

    ws["A1"] = note
    ws["A1"].font = Font(italic=True, color="595959", size=10)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columns), 2))
    ws["A1"].alignment = Alignment(vertical="center", wrap_text=True)

    for i, (header, width, required, help_text, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        cell = ws.cell(row=2, column=i, value=header)
        cell.fill = REQ_FILL if required else HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
        if help_text:
            cell.comment = Comment(help_text, "Template")
            cell.comment.width = 320
            cell.comment.height = 110
        ws.column_dimensions[letter].width = width
        if fmt:
            for r in range(3, LAST_ROW + 1):
                ws.cell(row=r, column=i).number_format = fmt

    ws.row_dimensions[2].height = 32

    for i, value in enumerate(example, start=1):
        cell = ws.cell(row=3, column=i, value=value)
        cell.font = EXAMPLE_FONT
        cell.fill = EXAMPLE_FILL

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}2"
    return ws


def dropdown(ws, column_letter, formula, prompt):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    dv.error = prompt
    dv.errorTitle = "Not one of the allowed values"
    dv.promptTitle = "Pick from the list"
    dv.prompt = prompt
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}3:{column_letter}{LAST_ROW}")


wb = Workbook()
wb.remove(wb.active)

# --- README ----------------------------------------------------------------

readme = wb.create_sheet("README")
readme.sheet_properties.tabColor = "C00000"
readme.column_dimensions["A"].width = 118

project_label = f"{CTX['project'][0]} — {CTX['project'][1]}" if CTX["project"] else "(no project yet)"
ws_label = ", ".join(f"{c} = {n}" for c, n in CTX["workstreams"]) or "(none yet)"

LINES = [
    ("AutoDock — planning workbook", "title"),
    (f"Project: {project_label}", "sub"),
    (f"Workstreams: {ws_label}", "sub"),
    ("", ""),
    ("Fill the sheets in order. Each one depends on the names you used in the one before it.", "body"),
    ("", ""),
    ("1. Team          people who will own or review work. Skip anyone already listed on Reference.", "body"),
    ("2. Milestones    the checkpoints inside each workstream.", "body"),
    ("3. Deliverables  optional. Named outputs under a milestone. Skip the sheet entirely if you", "body"),
    ("                 do not need this level.", "body"),
    ("4. Tasks         the actual work. This is the sheet that matters most.", "body"),
    ("5. Supply        optional. Long-lead orders and supply risks, so they can be watched.", "body"),
    ("", ""),
    ("Rules", "head"),
    ("", ""),
    ("RED header   = required. The row is rejected without it.", "body"),
    ("BLUE header  = optional. Leave it blank if you do not know yet.", "body"),
    ("", ""),
    ("Blank is a real answer. An empty Owner or Priority is loaded as \"not decided yet\" and shows up", "body"),
    ("on the clarification queue for you to resolve in the app. Nothing is ever guessed or filled in", "body"),
    ("on your behalf — so do not put a placeholder in just to fill the cell.", "body"),
    ("", ""),
    ("Row 3 of every sheet is a grey EXAMPLE row. It is ignored on import. Delete it or type over it.", "body"),
    ("", ""),
    ("Names are the link between sheets. A task finds its milestone by typing the milestone name", "body"),
    ("exactly as written on the Milestones sheet, for the same workstream. Dropdowns are wired up", "body"),
    ("where possible, so prefer picking over typing.", "body"),
    ("", ""),
    ("Dates: YYYY-MM-DD (e.g. 2026-10-20).", "body"),
    ("", ""),
    ("The one constraint to know about", "head"),
    ("", ""),
    (f"A task must have a start date and a due date, and may span at most {MAX_SPAN} calendar days.", "body"),
    (f"If a task covers more than one day, every day needs its own outcome (the Day 1..{MAX_SPAN} columns),", "body"),
    ("so a multi-day task is several checkable outcomes rather than one vague block of time.", "body"),
    ("", ""),
    ("This means long work has to be broken down before it can be entered. If that does not suit how", "body"),
    ("you want to plan, say so — the limit is one setting (MAX_TASK_SPAN_DAYS) and the requirement for", "body"),
    ("dates at all is a few lines of code. Both are easy to change; neither should be worked around by", "body"),
    ("putting fake dates in.", "body"),
    ("", ""),
    ("When you are done", "head"),
    ("", ""),
    ("Save the file and hand it back. It gets loaded in one pass: team first, then milestones,", "body"),
    ("deliverables, tasks, supply. You will get a report of what was created and anything that could", "body"),
    ("not be resolved. Nothing is written until the whole file validates.", "body"),
]

row = 1
for text, kind in LINES:
    cell = readme.cell(row=row, column=1, value=text)
    if kind == "title":
        cell.font = TITLE_FONT
    elif kind == "sub":
        cell.font = Font(bold=True, color="595959")
    elif kind == "head":
        cell.font = Font(bold=True, size=12, color="1F3864")
    else:
        cell.font = Font(size=10.5)
    row += 1

# --- Team ------------------------------------------------------------------

team = write_sheet(
    wb,
    "Team",
    [
        ("Name", 26, True, "Full name as it should appear in the app.", None),
        ("Email", 32, True, "Must be unique. This is how tasks on the Tasks sheet point at an owner.", None),
        ("Role", 14, True, "manager = can create tasks, set deadlines, manage the team.\nemployee = works on and updates their own tasks.", None),
        ("Title", 26, False, "Optional job title, shown on the roster.", None),
    ],
    ["EXAMPLE - delete this row", "example@vegam.co", "employee", "Mechanical Engineer"],
    "People to add. Anyone already in the system is listed on the Reference sheet — do not repeat them here.",
)
dropdown(team, "C", '"manager,employee"', "manager or employee")

# --- Milestones ------------------------------------------------------------

miles = write_sheet(
    wb,
    "Milestones",
    [
        ("Workstream", 14, True, "Code from the Reference sheet: " + ", ".join(WS_CODES), None),
        ("Milestone Name", 38, True, "Must be unique within the workstream. Tasks and deliverables refer back to this exact text.", None),
        ("Description", 46, False, "What 'done' means for this milestone.", None),
        ("Start Date", 14, False, "YYYY-MM-DD. Optional.", "yyyy-mm-dd"),
        ("End Date", 14, False, "YYYY-MM-DD. Optional.", "yyyy-mm-dd"),
        ("Target Date", 14, False, "The date being committed to, if different from End Date.", "yyyy-mm-dd"),
        ("Timeline Label", 16, False, "Free text for a relative schedule, e.g. 'W1 - W3'. Kept verbatim.", None),
        ("Long-Lead Items / Risks", 40, False, "Constraints and dependencies, one per line (Alt+Enter). Anything with a date belongs on the Supply sheet instead.", None),
    ],
    ["EXAMPLE - delete this row", "Chassis frame complete", "Welded frame, squared and load-tested", None, None, None, "W1 - W3", "Aluminium extrusion - 3 week lead"],
    "Checkpoints inside each workstream. A milestone groups tasks; it is not itself a task.",
)
dropdown(miles, "A", '"' + ",".join(WS_CODES) + '"', "Workstream code")

# --- Deliverables ----------------------------------------------------------

deliv = write_sheet(
    wb,
    "Deliverables",
    [
        ("Workstream", 14, True, "Code from the Reference sheet.", None),
        ("Milestone Name", 38, True, "Must match a Milestone Name on the Milestones sheet, in the same workstream.", None),
        ("Deliverable Name", 38, True, "The named output.", None),
        ("Description", 46, False, "Optional detail.", None),
        ("Due Date", 14, False, "YYYY-MM-DD. Optional.", "yyyy-mm-dd"),
        ("Due Label", 16, False, "Free text, e.g. 'by W2'. Kept verbatim.", None),
        ("Critical", 12, False, "yes = flagged as a critical long-lead item.", None),
    ],
    ["EXAMPLE - delete this row", "Chassis frame complete", "Load test report", "Signed-off test sheet", None, "by W3", "no"],
    "OPTIONAL SHEET. Only use it if you want a layer between a milestone and its tasks. Leave empty otherwise.",
)
dropdown(deliv, "A", '"' + ",".join(WS_CODES) + '"', "Workstream code")
dropdown(deliv, "G", '"yes,no"', "yes or no")

# --- Tasks -----------------------------------------------------------------

day_cols = [
    (
        f"Day {i} Outcome",
        28,
        False,
        f"What day {i} of this task produces. Required if the task spans {i} or more days.",
        None,
    )
    for i in range(1, MAX_SPAN + 1)
]

tasks = write_sheet(
    wb,
    "Tasks",
    [
        ("Workstream", 14, True, "Code from the Reference sheet. This is the only required parent.", None),
        ("Milestone Name", 32, False, "Optional. Must match the Milestones sheet exactly, same workstream. Blank = task sits directly under the workstream.", None),
        ("Deliverable Name", 28, False, "Optional, and only if a Milestone Name is given.", None),
        ("Task Title", 44, True, "Short, specific, one piece of work.", None),
        ("Description", 48, False, "Detail, acceptance criteria, links.", None),
        ("Owner Email", 28, False, "The single accountable person. Blank = unassigned, raised as a clarification.", None),
        ("Reviewer Email", 28, False, "Optional. Who signs the work off.", None),
        ("Supporters (emails, comma-separated)", 34, False, "Optional helpers. The owner stays singular and accountable.", None),
        ("Priority", 12, False, "P0 = drop everything ... P3 = whenever. Blank = not decided, raised as a clarification.", None),
        ("Start Date", 14, True, "YYYY-MM-DD. Required.", "yyyy-mm-dd"),
        ("Due Date", 14, True, f"YYYY-MM-DD. Required, and at most {MAX_SPAN} calendar days from the start (inclusive).", "yyyy-mm-dd"),
        *day_cols,
        ("Next Action", 30, False, "The immediate next step, if known.", None),
        ("Depends On (task titles)", 34, False, "Titles of other tasks in this sheet, comma-separated. They do not have to appear above this row.", None),
    ],
    [
        "EXAMPLE - delete this row",
        "Chassis frame complete",
        None,
        "Cut and square extrusion for base frame",
        "All 8 members cut to drawing, ends faced",
        (CTX["team"][0][1] if CTX["team"] else "example@vegam.co"),
        None,
        None,
        "P1",
        None,
        None,
        *(["Cut all 8 members to length"] + [None] * (MAX_SPAN - 1)),
        "Confirm stock on hand",
        None,
    ],
    "The work itself. Every task needs a workstream, a title, and a start and due date.",
)
dropdown(tasks, "A", '"' + ",".join(WS_CODES) + '"', "Workstream code")
dropdown(tasks, "B", f"=Milestones!$B$3:$B${LAST_ROW}", "A milestone name from the Milestones sheet")
dropdown(tasks, "I", '"P0,P1,P2,P3"', "P0, P1, P2 or P3")
dropdown(tasks, "F", f"=Reference!$H$3:$H${LAST_ROW}", "An email from Team or Reference")
dropdown(tasks, "G", f"=Reference!$H$3:$H${LAST_ROW}", "An email from Team or Reference")

# --- Supply ----------------------------------------------------------------

supply = write_sheet(
    wb,
    "Supply",
    [
        ("Workstream", 14, True, "Code from the Reference sheet.", None),
        ("Milestone Name", 32, True, "The milestone this item belongs to. Supply items hang off a milestone.", None),
        ("Item", 36, True, "What is being ordered, or the risk being tracked.", None),
        ("Kind", 12, True, "order = something to place by a date.\nrisk = a constraint with no order to place.", None),
        ("Lead Time (weeks)", 16, False, "Whole weeks from order to arrival.", None),
        ("Order By", 14, False, "YYYY-MM-DD. Leave blank and it is worked back from Needed By minus the lead time.", "yyyy-mm-dd"),
        ("Needed By", 14, False, "YYYY-MM-DD. The latest it can arrive and still be usable. Without this, lateness cannot be judged.", "yyyy-mm-dd"),
        ("Status", 16, False, "not_ordered / ordered / received / not_needed. Blank = not_ordered.", None),
        ("Owner Email", 28, False, "Who is chasing it.", None),
        ("Critical", 12, False, "yes = schedule stops without it.", None),
        ("Note", 40, False, "Vendor, quote reference, anything useful.", None),
    ],
    ["EXAMPLE - delete this row", "Chassis frame complete", "VFD unit", "order", 8, None, None, "not_ordered", None, "yes", "Quote pending from vendor"],
    "OPTIONAL SHEET. Long-lead orders and supply risks. An order placed late arrives late and nothing else can move — this is what lets the system warn you in advance.",
)
dropdown(supply, "A", '"' + ",".join(WS_CODES) + '"', "Workstream code")
dropdown(supply, "B", f"=Milestones!$B$3:$B${LAST_ROW}", "A milestone name from the Milestones sheet")
dropdown(supply, "D", '"order,risk"', "order or risk")
dropdown(supply, "H", '"not_ordered,ordered,received,not_needed"', "A supply status")
dropdown(supply, "J", '"yes,no"', "yes or no")

# --- Reference -------------------------------------------------------------

ref = wb.create_sheet("Reference")
ref.sheet_properties.tabColor = "808080"
ref["A1"] = "Read-only. Generated from the live system — do not edit."
ref["A1"].font = Font(italic=True, color="595959")

blocks = [
    (1, "Workstream Code", "Workstream Name", [(c, n) for c, n in CTX["workstreams"]], 14, 38),
    (4, "Existing Member", "Email", [(n, e) for n, e, _, _ in CTX["team"]], 24, 30),
]
for col, h1, h2, rows, w1, w2 in blocks:
    for off, header in enumerate((h1, h2)):
        cell = ref.cell(row=2, column=col + off, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ref.column_dimensions[get_column_letter(col)].width = w1
    ref.column_dimensions[get_column_letter(col + 1)].width = w2
    for i, (a, b) in enumerate(rows, start=3):
        ref.cell(row=i, column=col, value=a)
        ref.cell(row=i, column=col + 1, value=b)

ref.cell(row=2, column=7, value="Role / Priority / Status values").fill = HEADER_FILL
ref.cell(row=2, column=7).font = HEADER_FONT
ref.column_dimensions["G"].width = 32
for i, v in enumerate(
    [
        "Roles: manager, employee",
        "Priority: P0, P1, P2, P3",
        "Supply kind: order, risk",
        "Supply status: not_ordered, ordered, received, not_needed",
        f"Max task span: {MAX_SPAN} calendar days",
        "Dates: YYYY-MM-DD",
    ],
    start=3,
):
    ref.cell(row=i, column=7, value=v)

# Column H: every email a task may point at — existing members plus whatever
# gets typed on the Team sheet. This is what the Owner/Reviewer dropdowns read.
ref.cell(row=2, column=8, value="All emails (owner dropdown)").fill = HEADER_FILL
ref.cell(row=2, column=8).font = HEADER_FONT
ref.column_dimensions["H"].width = 32
for i, (_, email, _, _) in enumerate(CTX["team"], start=3):
    ref.cell(row=i, column=8, value=email)
start = 3 + len(CTX["team"])
for i in range(start, LAST_ROW + 1):
    ref.cell(row=i, column=8, value=f"=IF(Team!B{i - start + 3}=\"\",\"\",Team!B{i - start + 3})")

wb._sheets = [
    wb["README"],
    wb["Team"],
    wb["Milestones"],
    wb["Deliverables"],
    wb["Tasks"],
    wb["Supply"],
    wb["Reference"],
]
wb.active = 0

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)}  (max task span {MAX_SPAN} days, "
      f"{len(WS_CODES)} workstreams, {len(CTX['team'])} existing members)")
